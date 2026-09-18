# -*- coding: utf-8 -*-
"""生成题库**模板**（只有表头与一行占位示例，不含任何真实题目）。

    python tools/make_bank_template.py [输出目录]      默认 templates/

产出：
    bank_template.json   结构化题库模板
    bank_template.csv    带表头（含标准考试模板的冗余列）
    bank_template.xlsx   同上，Excel 版
    bank_template.docx   段落式（Word 2007+）
    bank_template.doc    段落式（Word 97-2003，OLE2）
    bank_template.txt    纯文本段落式

为什么单独做一个模板生成器
--------------------------
仓库里**不放任何真实题库**。示例题目属于本地内容，放在被 .gitignore 排除的
``data/local/sample_questions.json``；仓库对外只提供这里的空模板，
使用者克隆后填自己的题即可。

占位内容刻意写成「【示例】…」这种一眼能认出来的形式，
避免有人把示例当成真题目直接用。
"""

from __future__ import annotations

import csv
import json
import sys
from pathlib import Path
from typing import Dict, List

ROOT = Path(__file__).resolve().parent.parent

# 唯一的一行占位示例。不要在这里放任何真实题目。
PLACEHOLDER: Dict[str, object] = {
    "qtype": "单选题",
    "question": "【示例】请把这一行替换成你的题干？",
    "options": ["【示例】选项内容 A", "【示例】选项内容 B",
                "【示例】选项内容 C", "【示例】选项内容 D"],
    "answer": "A",
    "analysis": "【示例】解析内容，可以留空。",
}

# 标准考试模板的列（后四列程序会读但会自动忽略）
TABLE_HEADERS = ["题型", "题目", "选项", "答案", "解析",
                 "答案乱序", "阅卷人", "阅卷类型", "难度设置"]

_OPTIONS: List[str] = list(PLACEHOLDER["options"])          # type: ignore[arg-type]


# ---------------------------------------------------------------- 各格式

def _write_json(out_dir: Path) -> Path:
    p = out_dir / "bank_template.json"
    p.write_text(json.dumps(
        {"version": 1,
         "note": "题库模板：把 questions 里的示例换成你自己的题目即可。"
                 "题型/选项/答案/解析/题干 都是可识别的字段名。",
         "questions": [PLACEHOLDER]},
        ensure_ascii=False, indent=2), encoding="utf-8")
    return p


def _write_csv(out_dir: Path) -> Path:
    p = out_dir / "bank_template.csv"
    opts = _OPTIONS + [""] * (4 - len(_OPTIONS))
    with p.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["题型", "题目", "选项A", "选项B", "选项C", "选项D",
                    "答案", "解析", "难度设置"])
        w.writerow([PLACEHOLDER["qtype"], PLACEHOLDER["question"],
                    opts[0], opts[1], opts[2], opts[3],
                    PLACEHOLDER["answer"], PLACEHOLDER["analysis"], "适中"])
    return p


def _write_xlsx(out_dir: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "题库"
    ws.append(TABLE_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="EAF1FF")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    option_text = "\n".join(f"{'ABCDEFGH'[i]}. {o}" for i, o in enumerate(_OPTIONS))
    ws.append([PLACEHOLDER["qtype"], PLACEHOLDER["question"], option_text,
               PLACEHOLDER["answer"], PLACEHOLDER["analysis"],
               "否", "张三", "自动", "适中"])

    for i, width in enumerate([10, 56, 40, 10, 60, 10, 10, 10, 10], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    p = out_dir / "bank_template.xlsx"
    wb.save(str(p))
    return p


def _paragraph_lines(style: str) -> List[str]:
    """段落式模板的正文行。style 取 'bracket'（【题型】）或 'colon'（题型：）。"""
    lines: List[str] = []
    if style == "bracket":
        lines.append(f"【题型】{PLACEHOLDER['qtype']}")
        lines.append(f"1、{PLACEHOLDER['question']}")
        for i, o in enumerate(_OPTIONS):
            lines.append(f"{'ABCDEFGH'[i]}、{o}")
        lines.append(f"【答案】{PLACEHOLDER['answer']}")
        lines.append(f"【解析】{PLACEHOLDER['analysis']}")
    else:
        lines.append(f"1. {PLACEHOLDER['question']}")
        for i, o in enumerate(_OPTIONS):
            lines.append(f"{'ABCDEFGH'[i]}. {o}")
        lines.append(f"题型：{PLACEHOLDER['qtype']}")
        lines.append(f"答案：{PLACEHOLDER['answer']}")
        lines.append(f"解析：{PLACEHOLDER['analysis']}")
    return lines


def _write_docx(out_dir: Path) -> Path:
    from docx import Document

    doc = Document()
    doc.add_heading("题库模板（段落式 · Word 2007+）", level=1)
    for line in _paragraph_lines("colon"):
        doc.add_paragraph(line)
    doc.add_paragraph("")

    p = out_dir / "bank_template.docx"
    doc.save(str(p))
    return p


def _write_doc(out_dir: Path) -> Path:
    try:
        from tools.ole2_writer import write_word97_doc
    except ImportError:                       # 直接运行脚本时 sys.path[0] 是 tools/
        from ole2_writer import write_word97_doc  # type: ignore[no-redef]

    body = ["题库模板（段落式 · Word 97-2003）"] + _paragraph_lines("bracket") + [""]
    return write_word97_doc(out_dir / "bank_template.doc", "\n".join(body))


def _write_txt(out_dir: Path) -> Path:
    lines = ["题库模板（纯文本段落式）", ""]
    for line in _paragraph_lines("colon"):
        lines.append(line)
    p = out_dir / "bank_template.txt"
    p.write_text("\n".join(lines), encoding="utf-8")
    return p


# ---------------------------------------------------------------- 入口

WRITERS = [
    ("json", _write_json),
    ("csv", _write_csv),
    ("xlsx", _write_xlsx),
    ("docx", _write_docx),
    ("doc", _write_doc),
    ("txt", _write_txt),
]


def write_templates(out_dir: Path) -> List[Path]:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    produced: List[Path] = []
    for name, fn in WRITERS:
        try:
            produced.append(fn(out_dir))
        except Exception as exc:              # openpyxl / python-docx 缺失时跳过
            print(f"[跳过] 生成 {name} 模板失败：{exc}", file=sys.stderr)
    return produced


def main() -> int:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "templates"
    files = write_templates(out)
    print(f"已生成 {len(files)} 个题库模板到 {out}")
    for f in files:
        print(f"  - {f.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
