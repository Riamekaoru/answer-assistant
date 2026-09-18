# -*- coding: utf-8 -*-
"""生成多格式示例题库，用于首次跑通流程与验证各类导入器。

**题目内容不在代码里**。示例题目属于本地内容，存放在
``data/local/sample_questions.json``（已被 .gitignore 排除，不随仓库发布）。
本模块只负责把题目序列化成六种格式。

若本地没有这个文件，``write_samples()`` 返回空列表 —— 调用方应改为生成
``templates/`` 里的空模板（见 ``tools/make_bank_template.py``）。

产出（有本地题目时）：
    sample_bank.json    结构化题库（推荐格式）
    sample_bank.csv     UTF-8 带表头
    sample_bank.xlsx    带标准考试模板列（含答案乱序/阅卷人等冗余列，验证忽略逻辑）
    sample_bank.docx    段落式题库（1. 题干 / A. 选项 / 答案：X / 解析：X）
    sample_bank.doc     Word 97-2003 二进制（【题型】/【答案】/【解析】中文括号标记）
    sample_bank.txt     纯文本段落式

.doc 用 ``tools/ole2_writer`` 直接按格式规范拼出来，不依赖 Word / WPS / LibreOffice。
它同时承担两个职责：
    1. 让 ``core.legacy_doc`` 的 OLE2 + FIB + Piece Table 解析有真实二进制的回归测试；
    2. 覆盖中文办公文档常见的 ``【题型】/【答案】/【解析】`` + ``、`` 分隔符写法。

命令行：
    python tools/make_sample_bank.py [输出目录]
"""

from __future__ import annotations

import csv
import json
import sys
from pathlib import Path
from typing import Dict, List

ROOT = Path(__file__).resolve().parent.parent

# 本地示例题目（被 .gitignore 排除）。仓库里不放任何真实/示例题目。
LOCAL_QUESTIONS = ROOT / "data" / "local" / "sample_questions.json"


def load_questions() -> List[Dict[str, object]]:
    """读取本地示例题目；文件不存在或格式不对时返回空列表。"""
    if not LOCAL_QUESTIONS.exists():
        return []
    try:
        data = json.loads(LOCAL_QUESTIONS.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        print(f"[警告] 读取 {LOCAL_QUESTIONS} 失败：{exc}", file=sys.stderr)
        return []
    if isinstance(data, dict):
        items = data.get("questions")
    else:
        items = data
    if not isinstance(items, list):
        return []
    return [q for q in items if isinstance(q, dict) and q.get("question")]



def load_topic() -> str:
    """本地示例题目的主题名（仅用于写进 JSON 的说明字段）。"""
    if not LOCAL_QUESTIONS.exists():
        return "示例题库"
    try:
        data = json.loads(LOCAL_QUESTIONS.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return "示例题库"
    topic = data.get("topic") if isinstance(data, dict) else None
    return str(topic) if topic else "示例题库"


def write_samples(out_dir: Path) -> List[Path]:
    """写出全部示例题库文件，返回生成的文件列表。

    本地没有 ``data/local/sample_questions.json`` 时返回空列表 ——
    题目不进仓库，调用方应改为生成 ``templates/`` 里的空模板。
    """
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    produced: List[Path] = []

    questions = load_questions()
    if not questions:
        print(f"[跳过] 本地没有示例题目（{LOCAL_QUESTIONS}），未生成任何文件",
              file=sys.stderr)
        return produced

    # ---------- JSON ----------
    json_path = out_dir / "sample_bank.json"
    json_path.write_text(json.dumps(
        {"version": 1, "topic": load_topic(),
         "questions": questions},
        ensure_ascii=False, indent=2), encoding="utf-8")
    produced.append(json_path)

    # ---------- CSV ----------
    csv_path = out_dir / "sample_bank.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["题型", "题目", "选项A", "选项B", "选项C", "选项D",
                         "答案", "解析", "难度设置"])
        for q in questions:
            opts = list(q.get("options") or [])  # type: ignore[arg-type]
            opts += [""] * (4 - len(opts))
            writer.writerow([q["qtype"], q["question"], opts[0], opts[1], opts[2],
                             opts[3], q["answer"], q["analysis"], "适中"])
    produced.append(csv_path)

    # ---------- XLSX（模拟标准考试模板，含冗余列） ----------
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "题库"
        headers = ["题型", "题目", "选项", "答案", "解析",
                   "答案乱序", "阅卷人", "阅卷类型", "难度设置"]
        ws.append(headers)
        head_fill = PatternFill("solid", fgColor="EAF1FF")
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for q in questions:
            opts = list(q.get("options") or [])  # type: ignore[arg-type]
            option_text = "\n".join(
                f"{'ABCDEFGH'[i]}. {o}" for i, o in enumerate(opts))
            ws.append([q["qtype"], q["question"], option_text, q["answer"],
                       q["analysis"], "否", "张三", "自动", "适中"])

        widths = [10, 56, 40, 10, 60, 10, 10, 10, 10]
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

        xlsx_path = out_dir / "sample_bank.xlsx"
        wb.save(str(xlsx_path))
        produced.append(xlsx_path)
    except Exception as exc:  # openpyxl 缺失时跳过
        print(f"[跳过] 生成 xlsx 失败：{exc}", file=sys.stderr)

    # ---------- DOCX ----------
    try:
        from docx import Document

        doc = Document()
        doc.add_heading("示例练习题库", level=1)
        for i, q in enumerate(questions, start=1):
            doc.add_paragraph(f"{i}. {q['question']}")
            for j, opt in enumerate(list(q.get("options") or [])):  # type: ignore[arg-type]
                doc.add_paragraph(f"{'ABCDEFGH'[j]}. {opt}")
            doc.add_paragraph(f"题型：{q['qtype']}")
            doc.add_paragraph(f"答案：{q['answer']}")
            doc.add_paragraph(f"解析：{q['analysis']}")
            doc.add_paragraph("")

        docx_path = out_dir / "sample_bank.docx"
        doc.save(str(docx_path))
        produced.append(docx_path)
    except Exception as exc:
        print(f"[跳过] 生成 docx 失败：{exc}", file=sys.stderr)

    # ---------- DOC（Word 97-2003，中文括号标记写法） ----------
    try:
        # 直接运行本脚本时 sys.path[0] 是 tools/，作为模块导入时是项目根目录
        try:
            from tools.ole2_writer import write_word97_doc
        except ImportError:
            from ole2_writer import write_word97_doc  # type: ignore[no-redef]

        doc_lines: List[str] = []
        for i, q in enumerate(questions, start=1):
            doc_lines.append(f"【题型】{q['qtype']}")
            doc_lines.append(f"{i}、{q['question']}")
            for j, opt in enumerate(list(q.get("options") or [])):  # type: ignore[arg-type]
                doc_lines.append(f"{'ABCDEFGH'[j]}、{opt}")
            doc_lines.append(f"【答案】{q['answer']}")
            doc_lines.append(f"【解析】{q['analysis']}")
            doc_lines.append("")

        doc_path = write_word97_doc(out_dir / "sample_bank.doc", "\n".join(doc_lines))
        produced.append(doc_path)
    except Exception as exc:
        print(f"[跳过] 生成 doc 失败：{exc}", file=sys.stderr)

    # ---------- TXT ----------
    txt_lines: List[str] = []
    for i, q in enumerate(questions, start=1):
        txt_lines.append(f"{i}. {q['question']}")
        for j, opt in enumerate(list(q.get("options") or [])):  # type: ignore[arg-type]
            txt_lines.append(f"{'ABCDEFGH'[j]}. {opt}")
        txt_lines.append(f"题型：{q['qtype']}")
        txt_lines.append(f"答案：{q['answer']}")
        txt_lines.append(f"解析：{q['analysis']}")
        txt_lines.append("")
    txt_path = out_dir / "sample_bank.txt"
    txt_path.write_text("\n".join(txt_lines), encoding="utf-8")
    produced.append(txt_path)

    return produced


def main() -> int:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "data" / "samples"
    files = write_samples(out)
    print(f"已生成 {len(files)} 个示例题库文件到 {out}")
    for f in files:
        print(f"  - {f.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
